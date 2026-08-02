import io
import random
import re
from dataclasses import dataclass
from .pdf_builder_usecase import ReportPdfBuilder
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Cm, Inches

plt.switch_backend("Agg")


@dataclass
class StudentPerformance:
    name: str
    average: float
    status: str

@dataclass
class Diagnosis:
    descriptor_priority: pd.DataFrame
    students: list[StudentPerformance]

    @property
    def critical_descriptors(self) -> list[str]:
        return [str(item) for item in self.descriptor_priority.index.tolist()[:5]]


@dataclass
class ReportArtifacts:
    docx: bytes
    pdf: bytes


class ReportBuilder:
    margin = Cm(1.2)

    def build(
        self,
        diagnosis: Diagnosis,
        questions: dict[str, list[dict]],
        interventions: dict[str, list[dict]],
    ) -> ReportArtifacts:
        chart = self._build_priority_chart(diagnosis)
        selected_questions = self._select_questions(diagnosis.critical_descriptors, questions)
        docx_bytes = self._build_docx(
            diagnosis,
            selected_questions,
            interventions,
            chart,
        )
        pdf_bytes = self._build_pdf(
            diagnosis,
            selected_questions,
            interventions,
            chart,
        )

        return ReportArtifacts(docx=docx_bytes, pdf=pdf_bytes)

    def _build_priority_chart(self, diagnosis: Diagnosis) -> bytes:
        chart_stream = io.BytesIO()
        plt.figure(figsize=(9, 4))
        diagnosis.descriptor_priority["IPP"].head(10).plot(
            kind="bar", color="#9b1c31", edgecolor="black"
        )
        plt.title("Habilidades Criticas - IPP")
        plt.ylabel("Indice de Prioridade Pedagogica")
        plt.tight_layout()
        plt.savefig(chart_stream, format="png")
        plt.close()
        chart_stream.seek(0)
        return chart_stream.getvalue()

    def _build_docx(
        self,
        diagnosis: Diagnosis,
        selected_questions: dict[str, dict | None],
        interventions: dict[str, list[dict]],
        chart: bytes,
    ) -> bytes:
        document = Document()
        for section in document.sections:
            section.top_margin = self.margin
            section.bottom_margin = self.margin
            section.left_margin = self.margin
            section.right_margin = self.margin

        document.add_heading("RELATORIO PEDAGOGICO SAEB", 0).alignment = 1
        document.add_heading("1. Indice de Prioridade Pedagogica (IPP)", 1)
        document.add_picture(io.BytesIO(chart), width=Inches(6))

        document.add_heading("2. Desempenho por Estudante", 1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        header_cells = table.rows[0].cells
        header_cells[0].text = "STATUS"
        header_cells[1].text = "ESTUDANTE"
        header_cells[2].text = "MEDIA %"

        for student in sorted(diagnosis.students, key=lambda item: item.average, reverse=True):
            row_cells = table.add_row().cells
            row_cells[0].text = student.status
            row_cells[1].text = student.name
            row_cells[2].text = f"{student.average}%"

        self._add_questions(document, diagnosis.critical_descriptors, selected_questions)
        self._add_interventions(document, diagnosis.critical_descriptors, interventions)

        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def _add_questions(
        self,
        document: Document,
        descriptors: list[str],
        selected_questions: dict[str, dict | None],
    ) -> None:
        document.add_page_break()
        document.add_heading("3. Simulado de Reforco", 1)

        question_number = 1
        for descriptor in descriptors:
            question = selected_questions.get(descriptor.upper())
            if not question:
                document.add_paragraph(f"Nenhuma questao cadastrada para o descritor {descriptor}.")
                continue

            paragraph = document.add_paragraph()
            paragraph.add_run(f"Questao {question_number} (Descritor {descriptor}): ").bold = True
            paragraph.add_run(self._strip_image_tags(question.get("description", "")))

            for option in question.get("options", []):
                document.add_paragraph(str(option), style="List Bullet")

            question_number += 1

    def _add_interventions(
        self,
        document: Document,
        descriptors: list[str],
        interventions: dict[str, list[dict]],
    ) -> None:
        document.add_page_break()
        document.add_heading("4. Intervencoes de Robotica e IA", 1)

        for descriptor in descriptors:
            descriptor_interventions = interventions.get(descriptor.upper(), [])
            if not descriptor_interventions:
                run = document.add_paragraph().add_run(
                    f"Sugestao de robotica nao cadastrada para o descritor {descriptor}."
                )
                run.italic = True
                continue

            document.add_heading(f"Estrategia para {descriptor}", 2)
            for intervention in descriptor_interventions:
                if intervention.get("skill"):
                    document.add_paragraph(f"Habilidade: {intervention.get('skill')}")

                for detail in intervention.get("intervention_data", []):
                    title = detail.get("title", "Sem titulo")
                    paragraph = document.add_paragraph()
                    paragraph.add_run(title).bold = True
                    document.add_paragraph(f"Desafio: {detail.get('challenge', 'N/A')}")
                    document.add_paragraph(f"IA: {detail.get('integration', 'N/A')}")

    def _build_pdf(
        self,
        diagnosis: Diagnosis,
        selected_questions: dict[str, dict | None],
        interventions: dict[str, list[dict]],
        chart: bytes,
    ) -> bytes:
        return ReportPdfBuilder().build(
            diagnosis,
            diagnosis.critical_descriptors,
            selected_questions,
            interventions,
            chart,
            self._strip_image_tags,
        )

    def _select_questions(
        self,
        descriptors: list[str],
        questions: dict[str, list[dict]],
    ) -> dict[str, dict | None]:
        selected_questions = {}
        for descriptor in descriptors:
            descriptor_questions = questions.get(descriptor.upper(), [])
            selected_questions[descriptor.upper()] = (
                random.choice(descriptor_questions) if descriptor_questions else None
            )
        return selected_questions

    def _strip_image_tags(self, value: str) -> str:
        return re.sub(r"\[(.*?)\]", "", str(value)).strip()
