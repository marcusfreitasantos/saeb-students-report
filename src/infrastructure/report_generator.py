from docx import Document

class ReportGenerator:
    def generate(self, students, diagnosis):
        doc = Document()
        doc.add_heading("SAEB Report", 0)

        for student in students:
            doc.add_paragraph(f"{student.name} - {student.average}% - {student.status()}")

        doc.save("report.docx")